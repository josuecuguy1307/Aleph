"""
verify_presentacion_cableada.py — VARA del tipo `presentacion` y su formato .pptx.
[Convergencia · superficie 7 · punto 3]

QUÉ MIDE, Y POR QUÉ ESTAS TRES COSAS Y NO OTRAS
------------------------------------------------
Sumar un tipo a EL vocabulario rompe tres cosas distintas, y las tres se midieron
ANTES de tocar nada. Esta vara es la que las deja medidas para siempre:

  R1 · EL EXPORT MIENTE MUDO. `formats` gobierna `artifact_export`, y su `_generate`
       termina en «desconocido → markdown crudo». Con `pptx` declarado y sin
       generador, la descarga servía un .md llamado .pptx — sin 500 y sin aviso:

           pedido fmt='pptx'  ->  fmt SERVIDO='md'  mime='text/markdown'

       Por eso R1 no se conforma con «devuelve 200»: exige los BYTES de un OOXML
       real, con sus slides adentro. Un archivo que empieza en `PK` y no tiene
       `ppt/slides/` es un zip cualquiera, no una presentación.

  R2 · EL DIBUJO QUEDA MUDO. Los DOS registros de render (`render/render.js` y
       `render/sala-render.js`) calculan `coverage().uncovered`, que debe quedar vacío. Un
       canónico nuevo sin renderer ni delegación declarada cae a prosa en silencio.
       Se miden LOS DOS porque no se ven entre sí: durante esta obra el primero
       quedó verde y el segundo rojo, y sólo la sonda lo dijo.

  R3 · EL ESPEJO DRIFTEA. `render/vocabulary.js` se genera de `vocabulary.py`. Sin
       regenerar, el navegador ve una unión distinta a la del backend.

Y una cuarta que NO es rotura sino la razón de la obra:

  R4 · EL PUENTE. `presentation` es la palabra EXACTA que emite `claseDe()` de
       Oficina (`platform/workspaces/plugins/openwork.js:131`). Si el alias no
       aterriza en `presentacion`, el hueco que esta obra vino a tapar sigue ahí.

PROBADA CAYENDO (no basta con verde): cada bloque se corrió con su pieza rota y
las cuatro dieron rojo. Ver `# CAÍDA:` en cada sección.

    product/backend/.venv/bin/python qa/verify_presentacion_cableada.py
"""
from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

_AQUI = Path(__file__).resolve().parent
_RAIZ = _AQUI.parent
for _p in (str(_AQUI / "lib"), str(_RAIZ / "product" / "backend"), str(_RAIZ / "platform")):
    if _p not in sys.path:
        sys.path.insert(0, _p)

os.environ.setdefault("PUPPET_MOTOR_PERSISTE", "0")
# Hermético: la cache de descargas JAMÁS toca los datos reales.
_TMP = Path(tempfile.mkdtemp(prefix="vara-presentacion-"))
os.environ["ALEPH_DATA_DIR"] = str(_TMP / "data")

from artifacts import vocabulary as vocab            # noqa: E402
from app.phase1 import artifact_export as ex         # noqa: E402

_fallos = 0


def ok(cond, nombre, detalle=""):
    global _fallos
    if cond:
        print(f"  ✓ {nombre}")
    else:
        _fallos += 1
        print(f"  ✗ {nombre}" + (f"  →  {detalle}" if detalle else ""))
    return bool(cond)


def seccion(t):
    print(f"\n── {t} " + "─" * max(0, 74 - len(t)))


print("=" * 80)
print("VARA · `presentacion` Y SU .pptx  [Convergencia · superficie 7]")
print("=" * 80)
print(f"ALEPH_DATA_DIR = {os.environ['ALEPH_DATA_DIR']}")

# ── R0 · el tipo existe y declara lo que dijimos ──────────────────────────────
seccion("R0 · el sustantivo existe, y con los campos que se decidieron")

ok("presentacion" in vocab.CANONICAL, "`presentacion` es canónico")
ok(vocab.formats_for("presentacion") == ["pptx", "md"],
   "declara formats ['pptx','md'] — el default es el pptx, no el respaldo",
   str(vocab.formats_for("presentacion")))
ok(vocab.default_format("presentacion") == "pptx", "default_format = pptx")
# El `content` de este tipo es MARKDOWN (un encabezado = un slide), no un .pptx: un
# modelo lo escribe perfecto, igual que un `dashboard`. Lo que NO cruza acá es el
# binario que llega de un stack — ése es ficha, y lo mide verify_oficina_puente.
ok(vocab.TYPES["presentacion"]["producible_by_llm"] is True,
   "producible_by_llm TRUE — el contenido es markdown, y eso el modelo lo escribe")
ok("presentacion" in vocab.producible_by_llm(),
   "…y está en el conjunto que se le ofrece")
# El executor escanea `*.<tipo>.json` del workdir; un markdown de slides no es eso.
ok(vocab.TYPES["presentacion"]["rich_capture"] is False,
   "rich_capture FALSE — el executor no lo barre del workdir")
ok("presentacion" not in vocab.rich_capture_types(), "…y no está en el set de captura rica")

# ── R0.bis · el clasificador LA OFRECE DE VERDAD ─────────────────────────────
seccion("R0.bis · no basta con estar en la lista: tiene que llegar al prompt")
# CAÍDA, y es más dura de lo que uno esperaría: con `producible_by_llm: False` el
# módulo NO LLEGA A IMPORTARSE. Su propio assert de la línea 186 —«classifier alias
# points at a type the model may not produce»— revienta porque `slides`,
# `diapositivas` y las otras tres apuntan a un tipo que el modelo no podría producir.
# O sea que el vocabulario se defiende solo, antes de que ninguna vara corra: apagar
# `producible_by_llm` sin sacar sus palabras coloquiales es un estado imposible, no un
# estado silencioso. Medido en esta obra.
#
# Se mide el TEXTO del menú, no la tupla: la tupla es lo que el código cree y el menú
# es lo que el modelo lee. Un tipo puede estar en la primera y no llegar al segundo.

from app.phase1 import stream_chat as sc                  # noqa: E402

ok(set(vocab.producible_by_llm()) == set(sc._ARTIFACT_TYPES),
   "el espejo no drifteó: el conjunto del vocabulario ES el del clasificador",
   f"{sorted(set(vocab.producible_by_llm()) ^ set(sc._ARTIFACT_TYPES))}")
_menu = sc._type_menu(sc._ARTIFACT_TYPES)
ok('"presentacion"' in _menu, "`presentacion` aparece en el menú textual del prompt")
# Y con GLOSA propia, no sólo con su `label_es`: el modelo tiene que saber CÓMO se
# estructura (encabezado = slide), o escribiría prosa corrida y saldría un solo slide.
ok("cada encabezado abre un slide" in _menu,
   "…y su glosa explica que un encabezado abre un slide")
# Las palabras con las que un humano lo pide, tal como el modelo las devuelve.
for _w in ("presentación", "slides", "diapositivas", "powerpoint", "keynote", "deck"):
    ok(sc._coerce_type(_w) == "presentacion", f"el clasificador acepta «{_w}»",
       str(sc._coerce_type(_w)))

# ── R1 · el export entrega un pptx DE VERDAD ─────────────────────────────────
seccion("R1 · el .pptx son bytes de OOXML con slides, no un markdown disfrazado")
# CAÍDA: comentando la rama `if fmt == "pptx"` de `_generate`, los cuatro checks de
# este bloque dan rojo (fmt servido = 'md', magic != PK, 0 slides, mime markdown).

MD = "# Propuesta Q3\n\n- Crecer 20%\n- Abrir dos mercados\n\n## Riesgos\n\n| eje | nivel |\n| --- | --- |\n| costo | alto |\n"
art = {"id": "art-vara", "type": "presentacion", "title": "Pitch", "content": MD}
info = ex.export(art, "pptx", "sid-vara")

ok(info["fmt"] == "pptx", "el formato SERVIDO es pptx (no cayó al markdown crudo)", info["fmt"])
ok(info["filename"].endswith(".pptx"), "…y el archivo se llama .pptx", info["filename"])
ok(info["mime"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation",
   "…con el mime de presentationml", info["mime"])

_bytes = Path(info["path"]).read_bytes()
ok(_bytes[:2] == b"PK", "los bytes son un contenedor OOXML (magic PK)", repr(_bytes[:4]))
_slides = []
try:
    _slides = [n for n in zipfile.ZipFile(io.BytesIO(_bytes)).namelist()
               if n.startswith("ppt/slides/slide")]
except Exception as exc:                                    # noqa: BLE001
    ok(False, "el zip se abre", str(exc))
ok(len(_slides) == 2, "dos encabezados → dos slides", f"{len(_slides)} slides")

# El corte por encabezado, leído del archivo y no del generador.
try:
    from pptx import Presentation
    _p = Presentation(io.BytesIO(_bytes))
    _titulos = [s.shapes.title.text for s in _p.slides]
    ok(_titulos == ["Propuesta Q3", "Riesgos"],
       "los títulos salen de los encabezados markdown, en orden", str(_titulos))
except ImportError:
    ok(False, "python-pptx importable — está DECLARADO en requirements.txt",
       "no se pudo importar `pptx`")

# NUNCA cero slides: un .pptx vacío no se abre. Es la misma regla que el puente
# («una ficha honesta antes que un marco en blanco»).
for _caso, _cont in (("sin encabezados", "una sola linea"), ("contenido vacío", "")):
    _b = ex._gen_pptx(_cont, "Obra")
    _n = len([n for n in zipfile.ZipFile(io.BytesIO(_b)).namelist()
              if n.startswith("ppt/slides/slide")])
    ok(_n == 1, f"borde «{_caso}» → 1 slide, jamás cero", f"{_n} slides")

# El respaldo declarado sigue existiendo: pedir md sobre el mismo tipo da md.
_md = ex.export({**art, "id": "art-vara-md"}, "md", "sid-vara")
ok(_md["fmt"] == "md", "el segundo formato declarado (md) sigue sirviendo markdown", _md["fmt"])

# ── R2 · los DOS registros de dibujo lo cubren ───────────────────────────────
seccion("R2 · cobertura de dibujo en LOS DOS registros (no se ven entre sí)")
# CAÍDA: sacando `presentacion` del DELEGATES de render/sala-render.js, este bloque da
# `uncovered=['presentacion']` en el registro `sala` y queda rojo. Medido: durante la
# obra pasó exactamente eso, con el otro registro ya verde.

_sonda = _AQUI / "lib" / "sonda_vocabulario.mjs"
_r = subprocess.run(["node", str(_sonda), str(_RAIZ)], capture_output=True, text=True, timeout=90)
if _r.returncode != 0:
    ok(False, "la sonda de registros corre", (_r.stderr or "")[:200])
else:
    _d = json.loads(_r.stdout)
    for _reg in ("aleph", "sala"):
        _c = _d.get(_reg) or {}
        ok(_c.get("uncovered") == [],
           f"registro `{_reg}`: ningún canónico sin dibujo ni delegación",
           str(_c.get("uncovered")))
        ok(_c.get("delegates", {}).get("presentacion") == "informe",
           f"registro `{_reg}`: `presentacion` delega a `informe`, DECLARADO",
           str(_c.get("delegates")))

# ── R3 · el espejo JS no driftea ─────────────────────────────────────────────
seccion("R3 · el navegador ve la misma unión que el backend")
# CAÍDA: con el tipo agregado y sin correr el generador, `--check` sale 1 con
# «DRIFT: … no coincide con vocabulary.py». Medido en esta obra antes de regenerar.

_g = subprocess.run([sys.executable, str(_RAIZ / "platform" / "artifacts" / "gen_vocabulary_js.py"),
                     "--check"], capture_output=True, text=True, timeout=60)
ok(_g.returncode == 0, "gen_vocabulary_js.py --check: el espejo coincide",
   (_g.stdout + _g.stderr).strip()[:200])
_espejo = (_RAIZ / "product" / "app" / "design" / "render" / "vocabulary.js").read_text(encoding="utf-8")
ok("presentacion" in _espejo, "…y el espejo lleva el tipo nuevo")

# ── R4 · el puente: la palabra de Oficina aterriza ───────────────────────────
seccion("R4 · `presentation` —la palabra EXACTA de Oficina— aterriza en el tipo")
# CAÍDA: sacando el alias `presentation` de ALIASES, normalize() devuelve None y el
# borde de escritura rechaza — que es el 422 que esta obra vino a sacar.

for _palabra in ("presentation", "pptx", "deck"):
    ok(vocab.normalize(_palabra) == "presentacion",
       f"`{_palabra}` → presentacion", str(vocab.normalize(_palabra)))

# La palabra viene del código de Oficina, no de la cabeza del que escribe la vara.
_plugin = (_RAIZ / "platform" / "workspaces" / "plugins" / "openwork.js").read_text(encoding="utf-8")
ok('return "presentation"' in _plugin,
   "y `presentation` es literalmente lo que openwork.js emite para un .pptx")

# Un alias jamás tapa un canónico (el assert del módulo, medido desde afuera).
ok(not (set(vocab.ALIASES) & set(vocab.CANONICAL)), "ningún alias nuevo tapa un canónico")

print()
print("=" * 80)
print("TODO VERDE" if _fallos == 0 else f"{_fallos} FALLO(S)")
print("=" * 80)
sys.exit(1 if _fallos else 0)
