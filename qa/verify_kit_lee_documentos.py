"""verify_kit_lee_documentos.py — ¿PUEDE EL KIT LEER UN PDF, UN DOCX, UN XLSX, UN CSV Y UN PNG?
[T2.5 · la pregunta que decide si adjuntar documentos es obra o cableado]

LA PRUEBA DE LA VISIÓN, APLICADA A DOCUMENTOS. Cada archivo lleva un SECRETO que no está
en ningún otro lado y que no se puede adivinar (ZANAHORIA-4417, TELESCOPIO-5518, …). Si el
secreto vuelve, se leyó. Si no vuelve, no se leyó — y no hay forma de fingirlo. Por eso la
vara no mira «devolvió 200» ni «no explotó»: mira el secreto.

CONTRA LOS SERVERS REALES DEL KIT, levantados con el cliente de PROD (`assembler.MCPServer`)
y con la config de `belt-kit.mcp.json` — no con una librería aparte que «hace lo mismo».

QUÉ CONTESTA, Y POR QUÉ IMPORTA. Oficina resuelve los adjuntos escribiendo el archivo en un
inbox del workspace y pasándole al agente una RUTA (`openwork/.../sync/attachment-file-part.ts`).
Esta vara mide si ese camino ya está adentro de la casa. Medido el 2026-08-19:

    pdf  · docx · xlsx  → markitdown.convert_to_markdown   LEÍDO
    csv                 → filesystem.read_text_file        LEÍDO
    png                 → markitdown                       NO (devuelve vacío: no hace OCR)

⇒ los DOCUMENTOS no necesitan visión ni código nuevo: sólo que el archivo llegue al workdir.
⇒ la IMAGEN sí necesita visión, y `cli.claude_cli`/`cli.codex_cli` NO la declaran
  (`centro_modelos._PICKER_HOSTEADO`). El PNG en rojo NO es un fallo de esta vara: es el
  control negativo que prueba que sabe dar rojo.

    product/backend/.venv/bin/python qa/verify_kit_lee_documentos.py
"""
from __future__ import annotations

import json
import os
import sys
import tempfile
from pathlib import Path

_REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_REPO / "platform" / "assembler"))
from assembler import MCPServer  # noqa: E402

WD = Path(tempfile.mkdtemp(prefix="e2e-lectura-"))
SECRETOS = {
    "pdf":  "ZANAHORIA-4417",
    "csv":  "BUFANDA-9032",
    "docx": "TELESCOPIO-5518",
    "xlsx": "CARDUMEN-7741",
    "png":  "MEDUSA-2260",
}


def hacer_archivos() -> dict:
    rutas = {}
    # PDF
    from reportlab.pdfgen import canvas as _c
    p = WD / "informe.pdf"
    cv = _c.Canvas(str(p)); cv.drawString(72, 720, f"Codigo de referencia: {SECRETOS['pdf']}"); cv.save()
    rutas["pdf"] = p
    # CSV
    p = WD / "datos.csv"
    p.write_text(f"clave,valor\nreferencia,{SECRETOS['csv']}\ntotal,42\n", encoding="utf-8")
    rutas["csv"] = p
    # DOCX
    import docx
    d = docx.Document(); d.add_paragraph(f"Codigo de referencia: {SECRETOS['docx']}")
    p = WD / "contrato.docx"; d.save(str(p)); rutas["docx"] = p
    # XLSX
    import openpyxl
    wb = openpyxl.Workbook(); wb.active["A1"] = "referencia"; wb.active["B1"] = SECRETOS["xlsx"]
    p = WD / "planilla.xlsx"; wb.save(str(p)); rutas["xlsx"] = p
    # PNG — el secreto DIBUJADO, que sólo se lee con visión (u OCR)
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (420, 90), "white")
    ImageDraw.Draw(img).text((10, 40), f"Codigo: {SECRETOS['png']}", fill="black")
    p = WD / "captura.png"; img.save(str(p)); rutas["png"] = p
    return rutas


def servidor(nombre: str) -> MCPServer:
    kit = json.loads((_REPO / "catalog/templates/kit/belt-kit.mcp.json").read_text())
    scfg = kit["mcpServers"][nombre]
    env = dict(os.environ)
    env["PUPPET_WORKDIR"] = str(WD)
    env["PUPPET_BELTS"] = str(_REPO / "product" / "belts")
    exp = lambda s: str(s).replace("${PUPPET_WORKDIR}", str(WD)).replace("${PUPPET_BELTS}", env["PUPPET_BELTS"])
    child = dict(env)
    for k, v in (scfg.get("env") or {}).items():
        child[k] = exp(v)
    return MCPServer(nombre, exp(scfg.get("command", "")),
                     [exp(a) for a in (scfg.get("args") or [])], env=child)


def main() -> int:
    rutas = hacer_archivos()
    print(f"\nworkdir: {WD}\n")
    print(f"{'formato':<8} {'tool':<28} {'secreto':<18} veredicto")
    print("-" * 78)

    mk = servidor("markitdown")
    fs = servidor("filesystem")
    ok_mk = mk.start()
    ok_fs = fs.start()

    filas = []
    for fmt in ("pdf", "docx", "xlsx", "csv", "png"):
        ruta, secreto = str(rutas[fmt]), SECRETOS[fmt]
        salida, tool = "", ""
        try:
            if fmt == "csv":
                tool = "filesystem.read_text_file"
                r = fs.call_tool("read_text_file", {"path": ruta}) if ok_fs else "server caído"
            else:
                tool = "markitdown.convert_to_markdown"
                r = mk.call_tool("convert_to_markdown", {"uri": Path(ruta).as_uri()}) if ok_mk else "server caído"
            salida = r if isinstance(r, str) else json.dumps(r, ensure_ascii=False)
        except Exception as e:
            salida = f"EXCEPCIÓN {type(e).__name__}: {e}"
        leyo = secreto in salida
        filas.append((fmt, leyo))
        print(f"{fmt:<8} {tool:<28} {secreto:<18} "
              f"{'LEÍDO' if leyo else 'NO — ' + salida[:44].replace(chr(10), ' ')}")

    for s in (mk, fs):
        try: s.stop()
        except Exception: pass

    print("-" * 78)
    leidos = [f for f, ok in filas if ok]
    print(f"leídos: {len(leidos)}/{len(filas)} → {', '.join(leidos) or 'ninguno'}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
