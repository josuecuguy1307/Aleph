"""
method_export.py — .aleph + exports de la pieza MÉTODO (ORDEN §7).

Retención por arquitectura, no por regla:
  · `.aleph` = extensión ÚNICA con `type` interno (agent|method). El archivo es
    AUTÓNOMO: un agente embebe los OBJETOS de sus métodos equipados (al importar
    entran a la biblioteca del receptor y se re-vinculan por referencia).
  · Lo ACUMULADO (run_count/last_run_at/adjust_log/procedencia de runs) NO viaja
    en NINGÚN export, ni premium — whitelist, no blacklist.
  · `keys` (byok_refs del dueño) se STRIPEA siempre: apuntan al vault del origen,
    en el receptor son basura peligrosa.
  · PDF simple = free. Word/Markdown/checklist/JSON crudo = premium (el muro lo
    aplica el router con tier_gate, fail-closed staged).

El PDF free es un writer mínimo hecho a mano (stdlib, Helvetica, Latin-1): la
casa no carga dependencias nuevas para un export simple.
"""

from __future__ import annotations

import json
from typing import Any, Optional

from app.phase1 import methods_repo as mr

ALEPH_VERSION = 1

# WHITELIST real (§7 pide "whitelist, no blacklist"): SOLO estos campos de contenido
# viajan. Todo lo demás — acumulado (run_count/last_run_at/adjust_log), server-owned
# (id/timestamps), y cualquier campo inesperado (incluido un 'keys'/'token' pegado por
# error o por un import hostil) — se DESCARTA. El paso conserva solo el schema del ORDEN §2.
_METHOD_CONTENT_KEYS = ("name", "steps", "phases", "requires")
_STEP_CONTENT_KEYS = ("id", "text", "phase", "checkpoint", "executor",
                      "evidence_hint", "timeout", "retries")


def _exec_humano(executor, es: bool = True) -> str:
    """[reforma · a] El executor referencia MCP→tool (`server#tool`). Eso es una coordenada;
    al PDF/Markdown que lee una persona va en llano, sin `#`."""
    s = str(executor or "").strip()
    if "#" not in s:
        return s
    srv, _, tool = s.partition("#")
    srv, tool = srv.strip(), tool.strip()
    if srv and tool:
        return f"{tool} ({'de ' if es else 'from '}{srv})"
    return tool or srv


class AlephError(Exception):
    def __init__(self, kind: str, detail: str = ""):
        super().__init__(detail or kind)
        self.kind = kind
        self.detail = detail or kind


def clean_method(method: dict) -> dict:
    """El objeto Method exportable: WHITELIST de campos de contenido (name/steps/
    phases; cada step solo el schema §2). Cero acumulado, cero server-keys, cero
    campos inesperados — un 'keys'/'token' colado no sobrevive a NINGÚN export."""
    m = method or {}
    out: dict = {k: m[k] for k in _METHOD_CONTENT_KEYS if k in m}
    steps = m.get("steps")
    if isinstance(steps, list):
        out["steps"] = [{k: s[k] for k in _STEP_CONTENT_KEYS if isinstance(s, dict) and k in s}
                        for s in steps]
    return out


def _strip_secrets(obj):
    """Deep-strip de credenciales de una receta antes de exportarla/importarla: borra
    `byok_ref` a CUALQUIER profundidad (vive en model.byok_ref, keys.<p>.byok_ref, …)
    + la clave top-level/anidada `keys`. Un byok_ref apunta al vault del ORIGEN: en el
    receptor es basura peligrosa (revela el UUID de cuenta del exportador). Whitelist
    sería ideal para la receta, pero la forma es nicho-agnóstica y abierta por diseño;
    el deny-list de campos-credencial es el mínimo hermético para el límite de exfil."""
    if isinstance(obj, dict):
        return {k: _strip_secrets(v) for k, v in obj.items()
                if str(k).lower() not in _SECRET_KEYS}
    if isinstance(obj, list):
        return [_strip_secrets(v) for v in obj]
    return obj


# claves-credencial que NUNCA viajan en un export (deny-list, case-insensitive)
_SECRET_KEYS = frozenset({
    "byok_ref", "keys", "api_key", "apikey", "token", "access_token", "refresh_token",
    "secret", "client_secret", "password", "passwd", "authorization", "auth", "bearer",
    "private_key", "cookie", "session", "credential", "credentials",
})


def method_to_aleph(method: dict) -> dict:
    return {
        "aleph_version": ALEPH_VERSION,
        "type": "method",
        "name": method.get("name") or "",
        "method": clean_method(method),
    }


def agent_to_aleph(puppet: dict, methods: list[dict]) -> dict:
    """El agente COMPLETO y autónomo: receta CON canvas (el usuario espera su
    Cuarto de vuelta — a diferencia del puente D3 que lo dropea), SIN `keys`,
    schema v1 forzado (v0/faltante = hijo silenciosamente no-resoluble)."""
    # deep-strip de credenciales (keys top-level Y byok_ref anidados en model.*, etc.)
    recipe = _strip_secrets(dict(puppet.get("config") or {}))
    recipe["schema_version"] = "v1"
    return {
        "aleph_version": ALEPH_VERSION,
        "type": "agent",
        "name": puppet.get("name") or "",
        "nicho": puppet.get("nicho") or "general",
        "recipe": recipe,
        # {ref: id-de-ORIGEN, method: objeto limpio} — el ref existe SOLO para
        # re-mapear belt.method_refs al importar (el receptor acuña ids NUEVOS;
        # el UUID de origen jamás es identidad del receptor)
        "methods": [{"ref": str(m.get("id") or ""), "method": clean_method(m)}
                    for m in methods],
    }


def parse_aleph(data: Any) -> tuple[str, dict]:
    """(type, payload) o AlephError tipado. El envoltorio va ALREDEDOR de la
    receta (nunca adentro: _ALLOWED_TOP_KEYS es cerrado)."""
    if not isinstance(data, dict):
        raise AlephError("aleph_invalid", "el .aleph no es un objeto JSON")
    try:
        version = int(data.get("aleph_version"))
    except (TypeError, ValueError, OverflowError):
        # OverflowError: int(Infinity) — el json de stdlib acepta Infinity/NaN por
        # default; un .aleph hostil no debe reventar el parser con un 500.
        raise AlephError("aleph_invalid", "falta o es inválido aleph_version")
    if version > ALEPH_VERSION:
        raise AlephError("aleph_version_unsupported",
                         f"aleph_version {version} > {ALEPH_VERSION} soportada")
    a_type = data.get("type")
    if a_type == "method":
        if not isinstance(data.get("method"), dict):
            raise AlephError("aleph_invalid", "type=method sin objeto method")
        return "method", data
    if a_type == "agent":
        if not isinstance(data.get("recipe"), dict):
            raise AlephError("aleph_invalid", "type=agent sin recipe")
        return "agent", data
    raise AlephError("aleph_invalid", f"type desconocido: {a_type!r}")


# ── renderers de export (el contenido; el MURO vive en el router) ─────────────

def _phase_groups(method: dict) -> list[tuple[str, list[dict]]]:
    groups: list[tuple[str, list[dict]]] = []
    for s in method.get("steps") or []:
        ph = (s.get("phase") or "").strip()
        if groups and groups[-1][0] == ph:
            groups[-1][1].append(s)
        else:
            groups.append((ph, [s]))
    return groups


def to_markdown(method: dict, lang: str = "es") -> str:
    es = lang != "en"
    out = [f"# {method.get('name', '')}", ""]
    n = 0
    for ph, steps in _phase_groups(method):
        if ph:
            out.append(f"## {ph}")
        for s in steps:
            n += 1
            mark = "🛑 " if s.get("checkpoint") else ""
            line = f"{n}. {mark}{s.get('text', '')}"
            hints = []
            if s.get("executor"):
                hints.append(("herramienta: " if es else "tool: ") + _exec_humano(s["executor"], es))
            if s.get("evidence_hint"):
                hints.append(("evidencia: " if es else "evidence: ") + str(s["evidence_hint"]))
            if hints:
                line += "  _(" + " · ".join(hints) + ")_"
            out.append(line)
        out.append("")
    if any(s.get("checkpoint") for s in method.get("steps") or []):
        out.append(("🛑 = checkpoint: el proceso espera tu OK." if es
                    else "🛑 = checkpoint: the process waits for your OK."))
    return "\n".join(out).strip() + "\n"


def to_checklist(method: dict, lang: str = "es") -> str:
    es = lang != "en"
    out = [method.get("name", ""), "=" * max(4, len(method.get("name", ""))), ""]
    for ph, steps in _phase_groups(method):
        if ph:
            out.append(f"{ph}:")
        for s in steps:
            mark = "[!]" if s.get("checkpoint") else "[ ]"
            out.append(f"  {mark} {s.get('text', '')}")
        out.append("")
    out.append(("[!] = requiere aprobación" if es else "[!] = requires approval"))
    return "\n".join(out).strip() + "\n"


def to_json_raw(method: dict) -> str:
    return json.dumps(clean_method(method), ensure_ascii=False, indent=2) + "\n"


def to_docx(method: dict, lang: str = "es") -> bytes:
    """Word/SOP formateado (python-docx, ya en el venv por rag_index)."""
    try:
        import io
        from docx import Document
    except Exception as e:
        raise AlephError("docx_backend_unavailable", type(e).__name__)
    es = lang != "en"
    doc = Document()
    doc.add_heading(method.get("name", ""), level=1)
    n = 0
    for ph, steps in _phase_groups(method):
        if ph:
            doc.add_heading(ph, level=2)
        for s in steps:
            n += 1
            mark = ("⛔ " if s.get("checkpoint") else "")
            p = doc.add_paragraph(f"{n}. {mark}{s.get('text', '')}")
            hints = []
            if s.get("executor"):
                hints.append(("herramienta: " if es else "tool: ") + _exec_humano(s["executor"], es))
            if s.get("evidence_hint"):
                hints.append(("evidencia esperada: " if es else "expected evidence: ")
                             + str(s["evidence_hint"]))
            if hints:
                r = p.add_run("  (" + " · ".join(hints) + ")")
                r.italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ── PDF simple (free) — writer mínimo stdlib ──────────────────────────────────

def _pdf_escape(s: str) -> str:
    return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def to_pdf(method: dict, lang: str = "es") -> bytes:
    es = lang != "en"
    lines: list[str] = [method.get("name", ""), ""]
    n = 0
    for ph, steps in _phase_groups(method):
        if ph:
            lines.append(f"{ph}:")
        for s in steps:
            n += 1
            mark = "[!] " if s.get("checkpoint") else ""
            lines.append(f"  {n}. {mark}{s.get('text', '')}")
    if any(s.get("checkpoint") for s in method.get("steps") or []):
        lines += ["", ("[!] = checkpoint: espera tu OK" if es
                       else "[!] = checkpoint: waits for your OK")]

    per_page = 46
    pages = [lines[i:i + per_page] for i in range(0, len(lines), per_page)] or [[]]

    objs: list[bytes] = []          # cuerpos, obj n = índice+1
    n_pages = len(pages)
    font_obj = 3 + 2 * n_pages
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n_pages))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")                     # 1
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())  # 2
    for i, page_lines in enumerate(pages):
        page_n = 3 + 2 * i
        stream_n = page_n + 1
        objs.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                     f"/Resources << /Font << /F1 {font_obj} 0 R >> >> "
                     f"/Contents {stream_n} 0 R >>").encode())
        content = ["BT /F1 11 Tf 56 750 Td 15 TL"]
        first = True
        for ln in page_lines:
            txt = _pdf_escape(ln)
            content.append(("" if first else "T* ") + "(" + txt + ") Tj")
            first = False
        content.append("ET")
        stream = "\n".join(content).encode("latin-1", errors="replace")
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
                    + stream + b"\nendstream")
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")  # font

    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_at}\n%%EOF\n").encode()
    return bytes(out)
