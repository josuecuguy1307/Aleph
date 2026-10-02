"""
rag_index.py — C1 · RAG (átomo Conocimiento) · MÓDULO DE INDEXACIÓN Y RECUPERACIÓN.

QUÉ RESUELVE
  El corpus del agente (docs que el usuario sube) vive en el Postgres LOCAL (la máquina
  del propio usuario en modo `self_hosted`). Este módulo es la MECÁNICA PURA — extracción
  de texto, chunking, embeddings BYOK, y el índice de $0: coseno en Python puro (stdlib
  `math`) sobre vectores guardados como arrays JSONB. NO numpy, NO pgvector, NO sqlite-vec
  (ausentes/frágiles en este venv por contrato). La DB no vive acá (el executor la maneja);
  este módulo solo transforma bytes → texto → chunks → vectores y puntúa por similitud.

INVARIANTES DE SEGURIDAD (revisión adversarial las caza)
  - La KEY DE EMBEDDINGS es del USUARIO (BYOK), se resuelve SERVER-SIDE, NUNCA se loguea,
    NUNCA se devuelve en HTTP. Sin key → error HONESTO (`resolve_embed_key` → None → el
    caller marca status='error_no_key'), JAMÁS la key de Aleph, JAMÁS un fallback silencioso.
  - Los chunks recuperados son CONTENIDO NO CONFIABLE del usuario: el assembler los enmarca
    como APUNTES, nunca como instrucciones; el piso de gates de dinero/envío es inmutable.
  - `embed_texts` es una función de módulo REEMPLAZABLE (monkeypatch): el harness offline la
    sustituye por un embedder determinista sin red ($0/committeable).

Stdlib + pypdf (6.14.2) + python-docx (1.2.0). pypdf/docx/repo/broker se importan PEREZOSO
(el import base del módulo es stdlib puro → barato y sin acoplar la DB).
"""

from __future__ import annotations

import hashlib
import io
import json
import math
import os
import re
import urllib.error
import urllib.request
from typing import Any, Optional

# ── EXCEPCIONES TIPADAS (el caller las mapea a status honesto) ──────────────────


class RagError(Exception):
    """Base de los errores del subsistema RAG."""


class RagIngestError(RagError):
    """Extracción de texto imposible (pdf escaneado sin OCR, archivo ilegible/cifrado, etc.).
    NUNCA indexamos basura: preferimos rechazar con razón honesta."""


class RagEmbedError(RagError):
    """Fallo al pedir embeddings (transporte/HTTP/parse/sin key). NUNCA arrastra la key."""


# ── DEFAULTS DE EMBEDDINGS por proveedor ────────────────────────────────────────
# base_url: endpoint OpenAI-compat SIN sufijo (embed_texts agrega '/embeddings').
#   openai → text-embedding-3-small (dim 1536).  Base: OPENAI_API_BASE | api.openai.com/v1.
#   gemini → gemini-embedding-001 (dim 3072) vía GEMINI_BASE (OpenAI-compat, /embeddings vivo).
#     [C1·verify-vivo] el endpoint OpenAI-compat de Gemini SOLO acepta gemini-embedding-001 hoy;
#     text-embedding-004 devuelve HTTP 404 ("not supported for embedContent"). Overrideable por env.
_OPENAI_BASE = os.environ.get("OPENAI_API_BASE", "https://api.openai.com/v1")
_GEMINI_BASE = os.environ.get(
    "GEMINI_API_BASE", "https://generativelanguage.googleapis.com/v1beta/openai"
)
_OPENAI_EMBED_MODEL = os.environ.get("PUPPET_EMBED_MODEL_OPENAI", "text-embedding-3-small")
_GEMINI_EMBED_MODEL = os.environ.get("PUPPET_EMBED_MODEL_GEMINI", "gemini-embedding-001")

EMBED_DEFAULTS: dict[str, dict[str, Any]] = {
    "openai": {"model": _OPENAI_EMBED_MODEL, "dim": 1536, "base_url": _OPENAI_BASE},
    "gemini": {"model": _GEMINI_EMBED_MODEL, "dim": 3072, "base_url": _GEMINI_BASE},
}

# Orden de preferencia al autodetectar proveedor (si la receta no fija embed_provider).
_PROVIDER_ORDER = ("openai", "gemini")

# Umbral del guard de PDF escaneado: menos alfanuméricos que esto ⇒ "sin texto" (raise).
_MIN_PDF_ALNUM = 8


def embed_defaults_for(provider: str) -> dict[str, Any]:
    """Devuelve {model, dim, base_url} para `provider` (o los de openai como piso).
    Copia defensiva (el caller puede pisar model/base_url sin mutar el registro)."""
    d = EMBED_DEFAULTS.get((provider or "").strip().lower()) or EMBED_DEFAULTS["openai"]
    return dict(d)


def embed_target_for(provider: str, model: Optional[str] = None) -> dict[str, Any]:
    """STEP 2·C1 · Finding #2 — resuelve {model, dim, base_url} para embeber SIGUIENDO un corpus YA
    indexado: parte de los defaults del `provider`, pero HONRA el `model` con el que el corpus quedó
    sellado (embed_model persistido) cuando difiere del default del provider. Así la CONSULTA se
    embebe con EXACTAMENTE el mismo modelo que los chunks — la recuperación sigue al corpus, no a la
    receta, y no hay drift silencioso. El base_url sale del provider (el endpoint OpenAI-compat no
    cambia por modelo)."""
    d = embed_defaults_for(provider)
    m = (model or "").strip()
    if m:
        d["model"] = m
    return d


# ── EXTRACCIÓN DE TEXTO (txt/md/docx/pdf) + guard anti-basura ───────────────────


def _alnum_count(s: str) -> int:
    return sum(1 for ch in s if ch.isalnum())


def _kind(name: Optional[str], mime: Optional[str]) -> str:
    """Clasifica el archivo en 'pdf' | 'docx' | 'text' por mime y/o extensión."""
    m = (mime or "").lower()
    ext = os.path.splitext(name or "")[1].lower().lstrip(".")
    if "pdf" in m or ext == "pdf":
        return "pdf"
    if "wordprocessingml" in m or "docx" in m or ext == "docx":
        return "docx"
    # markdown, text/*, sin extensión, o desconocido → tratamos como texto (best-effort).
    return "text"


def _as_text(data: Any) -> str:
    """Decodifica bytes a str de forma tolerante; passthrough si ya es str."""
    if isinstance(data, str):
        return data
    if isinstance(data, (bytes, bytearray)):
        b = bytes(data)
        try:
            return b.decode("utf-8")
        except UnicodeDecodeError:
            return b.decode("utf-8", errors="replace")
    return str(data or "")


def _to_bytes(data: Any) -> bytes:
    if isinstance(data, (bytes, bytearray)):
        return bytes(data)
    if isinstance(data, str):
        return data.encode("utf-8")
    raise RagIngestError("empty_or_invalid_bytes")


def _extract_docx(data: Any) -> str:
    try:
        from docx import Document  # lazy: solo si hay un docx
    except Exception as e:  # pragma: no cover - dependencia confirmada en el venv
        raise RagIngestError(f"docx_backend_unavailable: {type(e).__name__}")
    try:
        doc = Document(io.BytesIO(_to_bytes(data)))
    except Exception as e:
        raise RagIngestError(f"docx_unreadable: {type(e).__name__}")
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _extract_pdf(data: Any) -> str:
    try:
        from pypdf import PdfReader  # lazy
    except Exception as e:  # pragma: no cover - dependencia confirmada en el venv
        raise RagIngestError(f"pdf_backend_unavailable: {type(e).__name__}")
    try:
        reader = PdfReader(io.BytesIO(_to_bytes(data)))
    except Exception as e:
        raise RagIngestError(f"pdf_unreadable: {type(e).__name__}")
    # PDF cifrado sin password → no forzamos; señal honesta.
    if getattr(reader, "is_encrypted", False):
        try:
            if reader.decrypt("") == 0:  # 0 = no se pudo desencriptar con "" (pypdf)
                raise RagIngestError("pdf_encrypted")
        except RagIngestError:
            raise
        except Exception:
            raise RagIngestError("pdf_encrypted")
    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(pages)
    # GUARD ANTI-ESCANEADO: casi cero alfanuméricos ⇒ es un PDF de imágenes (sin OCR).
    # Rechazamos en vez de indexar basura vacía/ruidosa.
    if _alnum_count(text) < _MIN_PDF_ALNUM:
        raise RagIngestError("pdf_no_text")
    return text


def extract_text(name: Optional[str], mime: Optional[str], data: Any) -> str:
    """
    Extrae texto plano de un documento subido.

    Args:
      name: nombre del archivo (para inferir extensión). Puede ser None.
      mime: mime declarado (para inferir tipo). Puede ser None.
      data: bytes (docx/pdf/txt binario) o str (txt/md ya-texto).

    Formatos: txt/md (passthrough decodificado) · docx (python-docx) · pdf (pypdf por página
    con guard anti-escaneado). Lanza `RagIngestError` si el documento no rinde texto real.
    """
    kind = _kind(name, mime)
    if kind == "docx":
        text = _extract_docx(data)
    elif kind == "pdf":
        text = _extract_pdf(data)
    else:
        text = _as_text(data)
    if not (text and text.strip()):
        raise RagIngestError("empty_document")
    return text


def sha256_text(text: str) -> str:
    """Hash de CONTENIDO (hex) — dedup/procedencia, estable a través del contenedor."""
    return hashlib.sha256((text or "").encode("utf-8")).hexdigest()


def sha256_bytes(data: Any) -> str:
    """Hash de los bytes crudos (hex)."""
    return hashlib.sha256(_to_bytes(data)).hexdigest()


# ── CHUNKING (fronteras de párrafo/oración, tamaño ~fijo + overlap) ─────────────

_SENT_END = re.compile(r"[.!?…](?=\s)")


def _best_break(text: str, start: int, end: int, size: int) -> int:
    """
    Elige el mejor punto de corte dentro de [start, end): prioriza salto de párrafo,
    luego salto de línea, luego fin de oración. No corta demasiado temprano (>= mitad
    del tamaño) para evitar chunks minúsculos. Devuelve el índice ABSOLUTO tras el corte.
    """
    lo = max(1, size // 2)  # offset mínimo relativo a `start`
    window = text[start:end]
    for sep in ("\n\n", "\n"):
        idx = window.rfind(sep)
        if idx != -1 and idx >= lo:
            return start + idx + len(sep)
    best = -1
    for m in _SENT_END.finditer(window):
        pos = m.end()
        if pos >= lo:
            best = pos
    if best != -1:
        return start + best
    return end


def chunk_text(text: str, size: int = 1000, overlap: int = 150) -> list[tuple[int, str, int]]:
    """
    Parte `text` en chunks de ~`size` chars con `overlap` de solapamiento, cortando en
    fronteras naturales cuando existen. Devuelve [(chunk_ix, content, char_start)].

    char_start = índice ABSOLUTO en `text` del primer char del content (post-strip),
    para procedencia. Garantiza progreso monótono (nunca cicla).
    """
    text = text or ""
    if not text.strip():
        return []
    if size <= 0:
        size = 1000
    if overlap < 0:
        overlap = 0
    if overlap >= size:
        overlap = size // 4
    n = len(text)
    out: list[tuple[int, str, int]] = []
    start = 0
    ix = 0
    while start < n:
        end = min(start + size, n)
        cut = end if end >= n else _best_break(text, start, end, size)
        raw = text[start:cut]
        lead = len(raw) - len(raw.lstrip())
        content = raw.strip()
        if content:
            out.append((ix, content, start + lead))
            ix += 1
        if cut >= n:
            break
        nxt = cut - overlap
        if nxt <= start:
            nxt = start + 1  # progreso garantizado (evita loop infinito)
        start = nxt
    return out


# ── EMBEDDINGS (BYOK) — clon urllib de assembler._chat → POST /embeddings ───────


def _embed_batch_size() -> int:
    try:
        return max(1, int(os.environ.get("PUPPET_RAG_EMBED_BATCH", "64")))
    except ValueError:
        return 64


def _http_timeout() -> float:
    try:
        return float(os.environ.get("PUPPET_HTTP_TIMEOUT", "60"))
    except ValueError:
        return 60.0


def embed_texts(
    texts: list[str],
    *,
    provider: str,
    model: str,
    api_key: str,
    base_url: str,
) -> list[list[float]]:
    """
    Pide embeddings para `texts` a un endpoint OpenAI-compat (`base_url` + '/embeddings').

    Clon de assembler._chat: urllib POST, Bearer, User-Agent EXPLÍCITO (PUPPET_HTTP_UA;
    varios proveedores banean el UA por defecto de urllib) y timeout ACOTADO
    (PUPPET_HTTP_TIMEOUT). Batchea `texts` para no exceder límites del proveedor y preserva
    el ORDEN por el campo `index` de la respuesta. La key NUNCA se loguea ni se formatea en
    ninguna excepción.

    Devuelve una lista de vectores (list[float]) alineada 1:1 con `texts`.
    Lanza `RagEmbedError` ante cualquier fallo (transporte/HTTP/parse/sin key).

    NOTA: es una función de MÓDULO reemplazable por monkeypatch (harness offline).
    """
    if not texts:
        return []
    if not api_key:
        # Piso duro de la Directiva #2: embeddings SIEMPRE con la key del usuario. Sin key
        # NO caemos a la de Aleph ni a un default silencioso — es un error honesto.
        raise RagEmbedError("no_api_key")
    if not base_url:
        raise RagEmbedError("no_base_url")

    endpoint = base_url.rstrip("/") + "/embeddings"
    headers = {
        "Content-Type": "application/json",
        "User-Agent": os.environ.get("PUPPET_HTTP_UA", "puppet-ai/1.0"),
        "Authorization": f"Bearer {api_key}",  # solo header, jamás a logs/errores
    }
    timeout = _http_timeout()
    batch = _embed_batch_size()

    vectors: list[list[float]] = []
    for i in range(0, len(texts), batch):
        window = texts[i : i + batch]
        payload = {"model": model, "input": window}
        data = json.dumps(payload).encode()
        req = urllib.request.Request(endpoint, data=data, headers=headers, method="POST")
        try:
            with urllib.request.urlopen(req, timeout=timeout) as resp:
                body = json.loads(resp.read().decode())
        except urllib.error.HTTPError as e:
            # El body del proveedor puede traer la key ENMASCARADA por el propio proveedor
            # (nunca nuestro cleartext, que solo viaja en el header). Acotamos por las dudas.
            try:
                detail = e.read().decode(errors="replace")[:300]
            except Exception:
                detail = ""
            raise RagEmbedError(f"http_{e.code}: {detail}")
        except (urllib.error.URLError, TimeoutError, OSError) as e:
            reason = getattr(e, "reason", None) or e
            raise RagEmbedError(f"transport: {reason}")
        except (ValueError, json.JSONDecodeError) as e:
            raise RagEmbedError(f"parse: {type(e).__name__}")

        rows = body.get("data") if isinstance(body, dict) else None
        if not isinstance(rows, list) or len(rows) != len(window):
            raise RagEmbedError("bad_response_shape")
        # Preservar orden por `index` cuando el proveedor lo declara.
        try:
            rows = sorted(rows, key=lambda r: r.get("index", 0))
        except Exception:
            pass
        for r in rows:
            emb = r.get("embedding") if isinstance(r, dict) else None
            if not isinstance(emb, list) or not emb:
                raise RagEmbedError("missing_embedding")
            vectors.append([float(x) for x in emb])

    return vectors


# ── ÍNDICE DE $0 — coseno en Python puro (stdlib math) ──────────────────────────


def cosine(a: list[float], b: list[float]) -> float:
    """Similitud coseno. 0.0 si dimensiones no coinciden o algún vector es nulo."""
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = 0.0
    na = 0.0
    nb = 0.0
    for x, y in zip(a, b):
        dot += x * y
        na += x * x
        nb += y * y
    if na == 0.0 or nb == 0.0:
        return 0.0
    return dot / math.sqrt(na * nb)


def top_k(query_vec: list[float], chunks: list[dict], k: int = 5) -> list[tuple[dict, float]]:
    """
    Puntúa cada chunk por coseno contra `query_vec` y devuelve los `k` mejores como
    [(chunk, score)] ordenados desc. Chunks con embedding ausente o de DIMENSIÓN distinta
    a la query se SALTAN (invariante: no mezclamos espacios de embeddings incompatibles).
    """
    if not query_vec or not chunks:
        return []
    qdim = len(query_vec)
    scored: list[tuple[dict, float]] = []
    for ch in chunks:
        emb = ch.get("embedding") if isinstance(ch, dict) else None
        if not isinstance(emb, list) or len(emb) != qdim:
            continue  # dim mismatch / faltante → fuera (drift guard)
        scored.append((ch, cosine(query_vec, emb)))
    scored.sort(key=lambda t: t[1], reverse=True)
    return scored[: max(0, k)]


# ── RESOLUCIÓN DE KEY BYOK (server-side, jamás a HTTP/logs) ─────────────────────


def resolve_embed_key(user_id: Optional[str], provider: str) -> Optional[str]:
    """
    Resuelve la key de embeddings del USUARIO para `provider` vía credential_broker
    (misma ruta con aislamiento cross-user y descifrado en memoria que el resto del BYOK).

    Devuelve el cleartext o None si no hay key (el broker devuelve "" ante ausencia/fallo;
    lo normalizamos a None para que el caller marque status='error_no_key' — nunca un skip
    silencioso, nunca la key de Aleph). El valor NUNCA se loguea acá.
    """
    if not user_id or not provider:
        return None
    try:
        from app.phase1 import credential_broker  # lazy: no acopla la DB al import base
    except Exception:
        return None
    resolver = credential_broker.make_user_resolver(user_id)
    secret = resolver((provider or "").strip().lower())
    return secret or None


def pick_embed_provider(
    user_id: Optional[str], preferred: Optional[str] = None
) -> Optional[str]:
    """
    Elige el proveedor de embeddings: `preferred` (de recipe.rag.embed_provider) si el
    usuario tiene esa key; si no, el primer proveedor embedding-capable con key disponible.
    Devuelve el nombre del proveedor o None si el usuario no tiene NINGUNA key de embeddings.
    """
    pref = (preferred or "").strip().lower()
    if pref and pref in EMBED_DEFAULTS and resolve_embed_key(user_id, pref):
        return pref
    for prov in _PROVIDER_ORDER:
        if resolve_embed_key(user_id, prov):
            return prov
    return None


__all__ = [
    "RagError",
    "RagIngestError",
    "RagEmbedError",
    "EMBED_DEFAULTS",
    "embed_defaults_for",
    "embed_target_for",
    "extract_text",
    "sha256_text",
    "sha256_bytes",
    "chunk_text",
    "embed_texts",
    "cosine",
    "top_k",
    "resolve_embed_key",
    "pick_embed_provider",
]
