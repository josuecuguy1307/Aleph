#!/usr/bin/env python
"""Helper del JUEZ — recalcula referencias contra los artifacts de una corrida.
NO asigna calidad (eso es contra rúbrica, humano). Uso: grade_helper.py <run_dir>
Imprime por celda completada los datos que la rúbrica pide verificar.
"""
import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

run = Path(sys.argv[1])
results = json.loads((run / "results.json").read_text())
ROOT = Path(__file__).resolve().parents[2]
FX = ROOT / "catalog/evals/finanzas-tasks/fixtures"


def t1(wd):
    out = {}
    wb = load_workbook(wd / "estado_resultados.xlsx")
    for n in wb.sheetnames:
        ws = wb[n]
        if "limpi" in n.lower():
            out["filas_limpias"] = ws.max_row - 1
        out.setdefault("formulas", {})[n] = sum(
            1 for r in ws.iter_rows() for c in r
            if isinstance(c.value, str) and c.value.startswith("="))
    ref = json.loads((ROOT / "catalog/evals/finanzas-tasks/T1-reference.json").read_text())
    out["ref"] = {"filas": ref["filas_unicas"], "neta": ref["utilidad_neta"],
                  "bruta": ref["utilidad_bruta"]}
    # valores cacheados de fórmulas (si el agente recalculó con una lib o pegó)
    wbv = load_workbook(wd / "estado_resultados.xlsx", data_only=True)
    vals = [c.value for ws in wbv.worksheets for r in ws.iter_rows() for c in r
            if isinstance(c.value, (int, float))]
    out["neta_presente_en_valores"] = any(
        abs(v - ref["utilidad_neta"]) < 0.5 for v in vals)
    return out


def t3(wd):
    out = {}
    df = pd.read_csv(FX / "T3-comps-fundamentals.csv")
    df["ev"] = df.precio_usd * df.acciones_diluidas_m + df.deuda_total_musd - df.caja_musd
    out["ref_ev_ebitda"] = dict(zip(df.ticker, (df.ev / df.ebitda_ltm_musd).round(2)))
    wb = load_workbook(wd / "comps.xlsx")
    out["formulas"] = [c.value for ws in wb.worksheets for r in ws.iter_rows()
                       for c in r if isinstance(c.value, str) and c.value.startswith("=")][:12]
    return out


def t4(wd):
    df = pd.read_csv(FX / "T4-macro-series.csv")
    cpi = df[df.serie == "CPIAUCSL"].sort_values("fecha").valor.values
    return {"ref_cpi_yoy": round((cpi[-1] / cpi[-13] - 1) * 100, 3),
            "ref_cpi_mom_saar": round(((cpi[-1] / cpi[-2]) ** 12 - 1) * 100, 3),
            "trampa_x12": round((cpi[-1] / cpi[-2] - 1) * 12 * 100, 3)}


def t5(wd):
    wb = load_workbook(wd / "varianzas.xlsx")
    names = {n.lower(): n for n in wb.sheetnames}
    v = next((names[n] for n in names if "varianza" in n), None)
    out = {"hojas": wb.sheetnames}
    if v:
        ws = wb[v]
        out["cond_format_reglas"] = len(list(ws.conditional_formatting))
        out["celdas_texto"] = [c.value for r in ws.iter_rows() for c in r
                               if isinstance(c.value, str)][:40]
    return out


CHECKS = {"T1": t1, "T3": t3, "T4": t4, "T5": t5}
for cell in results:
    if not cell["completo"]:
        continue
    wd = run / cell["model"] / cell["task"]
    fn = CHECKS.get(cell["task"])
    print(f"\n=== {cell['model']} × {cell['task']} ===")
    if fn is None:
        print("(T2: el checker ya recalcula la referencia FE — ver results.json)")
        continue
    try:
        print(json.dumps(fn(wd), ensure_ascii=False, indent=1, default=str)[:1500])
    except Exception as e:
        print("error:", e)
