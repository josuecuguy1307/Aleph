#!/usr/bin/env python
"""Checker T5 — exit 0 = COMPLETÓ. Uso: t5.py <workdir>"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

wd = Path(sys.argv[1])
out = {"artifact": False, "hoja_varianzas": False, "formulas_cruzadas": 0,
       "cond_format": 0, "etiquetas_FU": 0, "bridge": False}
try:
    wb = load_workbook(wd / "varianzas.xlsx")
    out["artifact"] = True
    names = {n.lower(): n for n in wb.sheetnames}
    vname = next((names[n] for n in names if "varianza" in n), None)
    out["hoja_varianzas"] = vname is not None
    if vname:
        ws = wb[vname]
        out["cond_format"] = len(list(ws.conditional_formatting))
        fu, fx, br = 0, 0, False
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith("=") and ("Presupuesto" in v or "Real" in v):
                        fx += 1
                    if v.strip() in ("F", "U", "Favorable", "Desfavorable"):
                        fu += 1
                    if "bridge" in v.lower():
                        br = True
        out.update(formulas_cruzadas=fx, etiquetas_FU=fu, bridge=br)
    ok = out["hoja_varianzas"] and out["formulas_cruzadas"] >= 5 and out["cond_format"] >= 1
except Exception as e:
    out["error"] = str(e); ok = False
print(json.dumps(out, ensure_ascii=False))
sys.exit(0 if ok else 1)
