#!/usr/bin/env python
"""Checker T3 — exit 0 = COMPLETÓ. Uso: t3.py <workdir>"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

wd = Path(sys.argv[1])
out = {"artifact": False, "formulas_vivas": 0, "median": False, "footnotes_filing": 0,
       "nm_para_perdida": False}
try:
    wb = load_workbook(wd / "comps.xlsx")
    out["artifact"] = True
    f, med, foot, nm = 0, False, 0, False
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    if v.startswith("="):
                        f += 1
                        if "MEDIAN" in v.upper():
                            med = True
                    if "10-K" in v or "10-Q" in v:
                        foot += 1
                    if "n.m" in v.lower():
                        nm = True
    out.update(formulas_vivas=f, median=med, footnotes_filing=foot, nm_para_perdida=nm)
    ok = f >= 5 and med
except Exception as e:
    out["error"] = str(e); ok = False
print(json.dumps(out, ensure_ascii=False))
sys.exit(0 if ok else 1)
