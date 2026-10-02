#!/usr/bin/env python
"""Checker T1 — exit 0 = COMPLETÓ. Uso: t1.py <workdir>"""
import json, sys
from pathlib import Path
from openpyxl import load_workbook

wd = Path(sys.argv[1])
out = {"artifact": False, "sheets": [], "formulas_vivas": 0, "filas_limpias": None}
try:
    p = wd / "estado_resultados.xlsx"
    wb = load_workbook(p)  # formulas as-is
    out["artifact"] = True
    out["sheets"] = wb.sheetnames
    f = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and "SUM" in c.value.upper():
                    f += 1
    out["formulas_vivas"] = f
    for name in wb.sheetnames:
        if "limpi" in name.lower():
            out["filas_limpias"] = wb[name].max_row - 1
    ref = json.loads((Path(__file__).resolve().parents[1] / "T1-reference.json").read_text())
    out["ref_filas"] = ref["filas_unicas"]
    out["ref_utilidad_neta"] = ref["utilidad_neta"]
    ok = len(wb.sheetnames) >= 2 and f >= 3
except Exception as e:
    out["error"] = str(e); ok = False
print(json.dumps(out, ensure_ascii=False))
sys.exit(0 if ok else 1)
